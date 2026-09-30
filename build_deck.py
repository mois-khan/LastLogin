# Builds the LastLogin pitch deck (Team Thrivers · Hackprix S3) in the brand style.
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ── Brand palette (from CLAUDE.md) ──
PAPER  = RGBColor(0xFB, 0xF7, 0xF0)
CARD   = RGBColor(0xFF, 0xFF, 0xFF)
INK    = RGBColor(0x2B, 0x2A, 0x28)
GRAPH  = RGBColor(0x55, 0x52, 0x4C)
EMBER  = RGBColor(0xC8, 0x74, 0x3A)
EMBER6 = RGBColor(0xA6, 0x56, 0x21)
SAGE   = RGBColor(0x6E, 0x8B, 0x74)
SAGE6  = RGBColor(0x4F, 0x6E, 0x57)
MIST   = RGBColor(0x9A, 0x93, 0x8A)
LINE   = RGBColor(0xE7, 0xE0, 0xD4)
PALE_SAGE  = RGBColor(0xED, 0xF1, 0xEC)
PALE_EMBER = RGBColor(0xF7, 0xEC, 0xE0)
CANDLE = RGBColor(0xFC, 0xFB, 0xF8)
WISP   = RGBColor(0xD9, 0xD2, 0xC8)

# System-safe fonts that match the brand vibe (warm serif / clean sans / mono).
F_DISP = "Georgia"
F_BODY = "Segoe UI"
F_MONO = "Consolas"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def slide():
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = PAPER
    return s


def box(s, l, t, w, h, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    return tf


def run(p, text, size, color, font=F_BODY, bold=False, italic=False, track=None):
    r = p.add_run(); r.text = text
    f = r.font
    f.size = Pt(size); f.bold = bold; f.italic = italic; f.name = font
    f.color.rgb = color
    if track is not None:
        r._r.get_or_add_rPr().set("spc", str(int(track * 100)))
    return r


def soft_edge(shape, inches):
    try:
        spPr = shape._element.spPr
        for el in spPr.findall(qn("a:effectLst")):
            spPr.remove(el)
        eff = spPr.makeelement(qn("a:effectLst"), {})
        eff.append(eff.makeelement(qn("a:softEdge"), {"rad": str(int(inches * 914400))}))
        spPr.append(eff)
    except Exception:
        pass


def no_shadow(sh):
    try:
        sh.shadow.inherit = False
    except Exception:
        pass


def card(s, l, t, w, h, fill=CARD, line=LINE, lw=1.0, radius=0.06):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(lw)
    no_shadow(sh)
    try:
        sh.adjustments[0] = radius
    except Exception:
        pass
    return sh


def candle(s, cx, top, scale=1.0, lit=True):
    cx = Inches(cx)
    bw, bh = Inches(0.62 * scale), Inches(1.15 * scale)
    body = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx - bw / 2, Inches(top + 0.55 * scale), bw, bh)
    body.fill.solid(); body.fill.fore_color.rgb = CANDLE
    body.line.color.rgb = RGBColor(0xD8, 0xD3, 0xCA); body.line.width = Pt(1.25)
    no_shadow(body)
    try:
        body.adjustments[0] = 0.16
    except Exception:
        pass
    ww, wh = Inches(0.03 * scale), Inches(0.13 * scale)
    wick = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx - ww / 2, Inches(top + 0.40 * scale), ww, wh)
    wick.fill.solid(); wick.fill.fore_color.rgb = INK; wick.line.fill.background(); no_shadow(wick)
    if lit:
        hw, hh = Inches(0.56 * scale), Inches(0.74 * scale)
        halo = s.shapes.add_shape(MSO_SHAPE.OVAL, cx - hw / 2, Inches(top - 0.02 * scale), hw, hh)
        halo.fill.solid(); halo.fill.fore_color.rgb = PALE_EMBER; halo.line.fill.background()
        no_shadow(halo); soft_edge(halo, 0.13 * scale)
        fw, fh = Inches(0.20 * scale), Inches(0.46 * scale)
        flame = s.shapes.add_shape(MSO_SHAPE.OVAL, cx - fw / 2, Inches(top + 0.10 * scale), fw, fh)
        flame.fill.solid(); flame.fill.fore_color.rgb = EMBER; flame.line.fill.background()
        no_shadow(flame); soft_edge(flame, 0.03 * scale)
        iw, ih = Inches(0.09 * scale), Inches(0.21 * scale)
        inner = s.shapes.add_shape(MSO_SHAPE.OVAL, cx - iw / 2, Inches(top + 0.21 * scale), iw, ih)
        inner.fill.solid(); inner.fill.fore_color.rgb = RGBColor(0xF6, 0xC9, 0x8A); inner.line.fill.background()
        no_shadow(inner)
    else:
        fw, fh = Inches(0.14 * scale), Inches(0.30 * scale)
        wisp = s.shapes.add_shape(MSO_SHAPE.OVAL, cx - fw / 2, Inches(top + 0.18 * scale), fw, fh)
        wisp.fill.solid(); wisp.fill.fore_color.rgb = WISP; wisp.line.fill.background()
        no_shadow(wisp); soft_edge(wisp, 0.05 * scale)


def eyebrow(s, text, x=0.9, y=0.62):
    tf = box(s, x, y, 11.4, 0.4)
    run(tf.paragraphs[0], text.upper(), 12.5, EMBER, F_BODY, bold=True, track=2.2)


def title(s, text, x=0.9, y=1.02, size=37, w=11.5, color=INK):
    tf = box(s, x, y, w, 1.5)
    run(tf.paragraphs[0], text, size, color, F_DISP)


def footer(s, idx):
    tf = box(s, 0.9, 7.02, 9, 0.35)
    p = tf.paragraphs[0]
    run(p, "LastLogin", 9, EMBER, F_BODY, bold=True)
    run(p, "   ·   Team Thrivers   ·   Hackprix S3", 9, MIST, F_BODY)
    tf2 = box(s, 11.6, 7.02, 0.83, 0.35)
    p2 = tf2.paragraphs[0]; p2.alignment = PP_ALIGN.RIGHT
    run(p2, f"{idx:02d} / 08", 9, MIST, F_MONO)


def bullets(s, l, t, w, items, size=16, gap=11, numbered=False, color=GRAPH):
    tf = box(s, l, t, w, 4.5)
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.12; p.space_after = Pt(gap)
        if numbered:
            run(p, f"{i+1:02d}", size, EMBER, F_MONO, bold=True); run(p, "   ", size, MIST, F_MONO)
        else:
            run(p, "•   ", size, EMBER, F_BODY, bold=True)
        if isinstance(it, tuple):
            run(p, it[0], size, INK, F_BODY, bold=True); run(p, it[1], size, color, F_BODY)
        else:
            run(p, it, size, color, F_BODY)
    return tf


# ─────────────────────────── SLIDE 1 — TITLE ───────────────────────────
s = slide()
card(s, 0, 0, 13.333, 0.16, fill=EMBER, line=None)  # thin ember rule at the very top
candle(s, 6.667, 0.55, scale=1.12, lit=True)
tf = box(s, 1, 2.72, 11.333, 0.5); tf.paragraphs[0].alignment = PP_ALIGN.CENTER
run(tf.paragraphs[0], "HACKPRIX S3   ·   TEAM THRIVERS", 13, MIST, F_BODY, bold=True, track=2.6)
tf = box(s, 1, 3.12, 11.333, 1.4); tf.paragraphs[0].alignment = PP_ALIGN.CENTER
run(tf.paragraphs[0], "LastLogin", 72, INK, F_DISP)
tf = box(s, 1, 4.5, 11.333, 0.6); tf.paragraphs[0].alignment = PP_ALIGN.CENTER
run(tf.paragraphs[0], "Your last login isn't goodbye.", 23, EMBER6, F_DISP, italic=True)
tf = box(s, 2.4, 5.2, 8.533, 0.9); tf.paragraphs[0].alignment = PP_ALIGN.CENTER
tf.paragraphs[0].line_spacing = 1.25
run(tf.paragraphs[0], "A proactive digital-estate manager that activates only after a verified death — your vault, your wishes, and your voice, for the people you love.", 14.5, GRAPH, F_BODY)
tf = box(s, 1, 6.5, 11.333, 0.4); tf.paragraphs[0].alignment = PP_ALIGN.CENTER
run(tf.paragraphs[0], "Gemini  ·  ElevenLabs  ·  Sarvam  ·  Ethereum  ·  MongoDB  ·  Vultr", 11, MIST, F_MONO)

# ─────────────────────────── SLIDE 2 — PROBLEM ───────────────────────────
s = slide()
eyebrow(s, "The problem")
title(s, "When we die, our digital life dies messy.")
bullets(s, 0.9, 2.55, 7.0, [
    "Accounts, photos, and passwords — locked away forever.",
    "Crypto with no keys and no heirs is lost for good.",
    "No trusted way to say a real goodbye.",
    "“Just in case” notes are scattered, stale, and unsafe.",
], size=16.5, gap=14)
card(s, 8.35, 2.45, 4.05, 3.05, fill=CARD, line=LINE)
tf = box(s, 8.75, 3.05, 3.25, 2.0, anchor=MSO_ANCHOR.MIDDLE)
for i, ln in enumerate(["No keys.", "No heirs.", "Gone."]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.line_spacing = 1.1
    run(p, ln, 30, INK if i < 2 else EMBER6, F_DISP)
footer(s, 2)

# ─────────────────────────── SLIDE 3 — SOLUTION ───────────────────────────
s = slide()
eyebrow(s, "The solution")
title(s, "Calm by design. Active only when it must be.")
bullets(s, 0.9, 2.55, 7.0, [
    "One secure vault for accounts, assets, and last words.",
    "Releases nothing until death is verified by people you trust.",
    "2-of-3 guardians + an AI-checked certificate flip the switch.",
    "Crypto inheritance settles on-chain — no lawyer, no lock-out.",
], size=16.5, gap=13)
card(s, 8.35, 2.45, 4.05, 3.2, fill=PALE_SAGE, line=None)
candle(s, 10.375, 2.7, scale=0.8, lit=False)
tf = box(s, 8.65, 4.35, 3.45, 1.1, anchor=MSO_ANCHOR.MIDDLE); tf.paragraphs[0].alignment = PP_ALIGN.CENTER
tf.paragraphs[0].line_spacing = 1.12
run(tf.paragraphs[0], "Nothing unlocks\nwhile you're alive.", 18, SAGE6, F_DISP)
footer(s, 3)

# ─────────────────────────── SLIDE 4 — HOW IT WORKS ───────────────────────────
s = slide()
eyebrow(s, "How it works")
title(s, "Five steps, one continuous flow.")
bullets(s, 0.9, 2.5, 11.4, [
    ("Fill the vault  ", "— an AI will-assistant interviews you.  (Gemini)"),
    ("Clone your voice  ", "— record once; translate to your family's language.  (ElevenLabs · Sarvam)"),
    ("Fund inheritance  ", "— name a beneficiary on a smart contract.  (Ethereum)"),
    ("Verify the passing  ", "— 2 guardians + an AI cert check flip ACTIVE → EXECUTING.  (Gemini Vision)"),
    ("Family unlocks  ", "— the dashboard plays the message and shows the transfer."),
], size=15.5, gap=13, numbered=True)
footer(s, 4)

# ─────────────────────────── SLIDE 5 — FEATURES ───────────────────────────
s = slide()
eyebrow(s, "What's inside")
title(s, "Built like an estate, not an app.")
bullets(s, 0.9, 2.55, 5.7, [
    ("Encrypted vault ", "— AES-256-GCM, on-chain fingerprint."),
    ("Guardian trigger ", "— 2-of-3, no single point of trust."),
    ("On-chain inheritance ", "— immutable and uncontestable."),
], size=15.5, gap=16)
bullets(s, 6.95, 2.55, 5.5, [
    ("AI will-assistant ", "— a chat becomes a structured will."),
    ("Final messages ", "— your voice, your family's language."),
    ("Talk-with-them ", "— a safety-railed AI remembrance."),
], size=15.5, gap=16)
footer(s, 5)

# ─────────────────────────── SLIDE 6 — EMOTIONAL CORE ───────────────────────────
s = slide()
eyebrow(s, "Why it matters")
title(s, "The first inheritance you can hear.")
bullets(s, 0.9, 2.55, 7.0, [
    "Not just passwords — presence. Your real voice, not text.",
    "Spoken in Hindi, Tamil, Telugu… in your own tone.",
    "Loved ones can still talk with a gentle AI you leave behind.",
    "Guard-railed: never leaks secrets, never asks for money.",
], size=16, gap=13)
card(s, 8.35, 2.45, 4.05, 3.2, fill=PALE_EMBER, line=None)
tf = box(s, 8.7, 3.05, 3.35, 1.8, anchor=MSO_ANCHOR.MIDDLE)
run(tf.paragraphs[0], "“Beta, main\nyahan hoon.”", 27, EMBER6, F_DISP)
tf = box(s, 8.7, 4.75, 3.35, 0.6)
run(tf.paragraphs[0], "— played back in their own cloned voice", 12, GRAPH, F_BODY, italic=True)
footer(s, 6)

# ─────────────────────────── SLIDE 7 — TECH & SPONSORS ───────────────────────────
s = slide()
eyebrow(s, "How we built it")
title(s, "Every sponsor earns its place.")
bullets(s, 0.9, 2.55, 5.7, [
    ("Gemini ", "— will-assistant, cert vision, drafting."),
    ("ElevenLabs ", "— the person's own cloned voice."),
    ("Sarvam ", "— India-first translation + speech."),
], size=15.5, gap=16)
bullets(s, 6.95, 2.55, 5.5, [
    ("Ethereum ", "— trustless inheritance + ledger."),
    ("MongoDB ", "— encrypted vault + messages."),
    ("Vultr · GitHub ", "— hosting + full build trail."),
], size=15.5, gap=16)
footer(s, 7)

# ─────────────────────────── SLIDE 8 — CLOSE ───────────────────────────
s = slide()
card(s, 0, 7.34, 13.333, 0.16, fill=EMBER, line=None)
candle(s, 6.667, 0.95, scale=1.0, lit=True)
tf = box(s, 1.3, 2.75, 10.733, 1.6); tf.paragraphs[0].alignment = PP_ALIGN.CENTER
tf.paragraphs[0].line_spacing = 1.18
run(tf.paragraphs[0], "Your last login isn't the end —\nit's the first thing you leave behind.", 30, INK, F_DISP)
tf = box(s, 1.8, 4.7, 9.733, 1.5); tf.paragraphs[0].alignment = PP_ALIGN.CENTER
for i, ln in enumerate([
    "Live demo: the contract flips on Etherscan, the family dashboard unlocks.",
    "Next: time-capsule messages, account-closure automation, a chat-with-them twin.",
    "Emotional + technical + defensible across every sponsor track.",
]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER; p.line_spacing = 1.0; p.space_after = Pt(7)
    run(p, ln, 14, GRAPH, F_BODY)
tf = box(s, 1, 6.45, 11.333, 0.4); tf.paragraphs[0].alignment = PP_ALIGN.CENTER
run(tf.paragraphs[0], "TEAM THRIVERS   ·   HACKPRIX S3", 12, MIST, F_BODY, bold=True, track=2.4)

# ── metadata + save ──
cp = prs.core_properties
cp.title = "LastLogin — Hackprix S3"
cp.author = "Team Thrivers"
cp.subject = "Proactive digital-estate manager"

out = r"c:\Projects\lastlogin\LastLogin-Thrivers-Hackprix-S3.pptx"
prs.save(out)
print("saved:", out, "| slides:", len(prs.slides._sldIdLst))
